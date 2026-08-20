"""导入导出工具"""
import io
from datetime import date, datetime
from typing import List, Tuple

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


def export_students_to_excel(students: List[dict]) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "学生信息"

    headers = ["学号", "姓名", "性别", "出生日期", "联系电话", "邮箱", "班级", "状态", "家庭地址", "备注"]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    center = Alignment(horizontal="center", vertical="center")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin_border

    for row_idx, stu in enumerate(students, 2):
        values = [
            stu.get("student_no", ""),
            stu.get("name", ""),
            stu.get("gender", ""),
            str(stu.get("birth_date", "")) if stu.get("birth_date") else "",
            stu.get("phone", ""),
            stu.get("email", ""),
            stu.get("class_name", ""),
            stu.get("status", ""),
            stu.get("address", ""),
            stu.get("remark", ""),
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for i, w in enumerate([10, 8, 6, 12, 15, 25, 12, 8, 30, 30], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


VALID_GENDERS = {"男", "女"}
VALID_STATUSES = {"在读", "休学", "退学", "毕业"}


def _parse_birth_date(value, row_no: int) -> Tuple[date, str]:
    """把 Excel 单元格值解析为 date 对象。

    返回 (date 或 None, 错误信息或 "")。支持 date/datetime 单元格与常见日期字符串。
    """
    if value is None or str(value).strip() == "":
        return None, ""
    if isinstance(value, datetime):
        return value.date(), ""
    if isinstance(value, date):
        return value, ""
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y-%m-%d %H:%M:%S", "%Y年%m月%d日"):
        try:
            return datetime.strptime(text, fmt).date(), ""
        except ValueError:
            continue
    return None, f"第{row_no}行出生日期格式无法识别: {text!r}"


def parse_from_excel(file_bytes: bytes) -> Tuple[List[dict], List[str]]:
    """解析学生 Excel,返回 (学生数据列表, 解析错误列表)。

    学生数据字段: student_no/name/gender/birth_date/phone/email/class_name/status/address/remark。
    解析错误(含行号)单独返回,由调用方决定是否落库。
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    ws = wb.active
    students: List[dict] = []
    errors: List[str] = []
    headers = None

    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        row_no = row_idx + 1
        if row_idx == 0:
            headers = [str(c).strip() if c else "" for c in row]
            continue
        if headers is None:
            continue
        values = [str(c).strip() if c else "" for c in row]
        if not any(values):
            continue
        mapping = {h: v for h, v in zip(headers, values)}
        stu = {
            "student_no": mapping.get("学号", ""),
            "name": mapping.get("姓名", ""),
            "gender": mapping.get("性别", "男"),
            "birth_date": mapping.get("出生日期", None),
            "phone": mapping.get("联系电话", ""),
            "email": mapping.get("邮箱", ""),
            "class_name": mapping.get("班级", ""),
            "status": mapping.get("状态", "在读"),
            "address": mapping.get("家庭地址", ""),
            "remark": mapping.get("备注", ""),
        }
        if not stu["student_no"] or not stu["name"]:
            errors.append(f"第{row_no}行缺少学号或姓名,已跳过")
            continue
        # 性别与状态枚举校验
        if stu["gender"] not in VALID_GENDERS:
            errors.append(f"第{row_no}行学号 {stu['student_no']}: 性别 {stu['gender']!r} 非法(仅允许 男/女)")
            continue
        if stu["status"] not in VALID_STATUSES:
            errors.append(f"第{row_no}行学号 {stu['student_no']}: 状态 {stu['status']!r} 非法(仅允许 在读/休学/退学/毕业)")
            continue
        # 日期解析: 必须转为 date 对象,否则 SQLAlchemy Date 字段无法写入
        birth_date, err = _parse_birth_date(stu["birth_date"], row_no)
        if err:
            errors.append(f"第{row_no}行学号 {stu['student_no']}: {err}")
            continue
        stu["birth_date"] = birth_date
        students.append(stu)
    return students, errors
